#!/usr/bin/env python3
"""
Convertit un export .xlsx (une ligne = une région) en data/<THEME>.json
pour le site.

Usage :
    python3 convert.py <fichier.xlsx> <NOM_THEME>

Suppose que la première ligne du fichier contient les en-têtes, que la
colonne A est "Thème", la colonne B est "Région", et que toutes les colonnes
suivantes sont les champs à afficher (dans l'ordre où elles apparaissent).
La couleur de fond de chaque cellule (surlignage manuel) est reprise telle
quelle et utilisée par le site pour colorer la valeur.

Nécessite : openpyxl (pip install openpyxl --break-system-packages)
"""
import sys
import json
import re
import unicodedata
import openpyxl

REGION_CODES = {
    "ile de france": "IDF",
    "normandie": "NOR",
    "hauts de france": "HDF",
    "grand est": "GES",
    "bretagne": "BRE",
    "pays de la loire": "PDL",
    "centre val de loire": "CVL",
    "bourgogne franche comte": "BFC",
    "nouvelle aquitaine": "NAQ",
    "auvergne rhone alpes": "ARA",
    "occitanie": "OCC",
    "provence alpes cote d azur": "PACACOR",
    "paca": "PACACOR",
    "corse": "PACACOR",
    "outre mer": "DROM",
    "drom com": "DROM",
}


def normalize(text):
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-zA-Z0-9]+", " ", text).strip().lower()
    return text


def slugify(text):
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", normalize(text))).strip("_")


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)

    xlsx_path, theme_name = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    headers = [c.value for c in ws[1]]
    field_cols = list(enumerate(headers))[2:]  # skip Thème, Région
    fields = [[slugify(h), h] for _, h in field_cols]

    out = {"_fields": fields}
    unmatched = []

    for row in ws.iter_rows(min_row=2, values_only=False):
        region_raw = row[1].value
        if not region_raw:
            continue
        code = REGION_CODES.get(normalize(region_raw))
        if not code:
            unmatched.append(region_raw)
            continue
        entry = {"label": region_raw}
        for col_idx, header in field_cols:
            cell = row[col_idx]
            color = None
            if cell.fill and cell.fill.patternType:
                rgb = cell.fill.fgColor.rgb
                if rgb and rgb != "00000000":
                    color = "#" + rgb[-6:]
            entry[slugify(header)] = {"text": cell.value if cell.value is not None else "Non renseigné", "color": color}
        out[code] = entry

    out_path = f"data/{theme_name}.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)

    print(f"Écrit : {out_path} ({len(out) - 1} régions)")
    if unmatched:
        print("Régions non reconnues (à vérifier / ajouter à REGION_CODES) :", unmatched)


if __name__ == "__main__":
    main()
